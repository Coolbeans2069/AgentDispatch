"""extract plain text from a .docx Word document."""
import json, os, sys


def main(argv):
    path = argv[0] if argv else ""
    if not path or not os.path.isfile(path):
        print(json.dumps(dict(ok=False, error="give a .docx file")))
        return 2
    try:
        import docx
    except ImportError as e:
        print(json.dumps(dict(ok=False, error="python-docx missing: %s" % e)))
        return 3
    try:
        d = docx.Document(path)
        paras = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
        tables = []
        for ti, t in enumerate(d.tables[:20]):
            rows = []
            for row in t.rows[:40]:
                rows.append([c.text.strip()[:200] for c in row.cells])
            tables.append(dict(index=ti, rows=rows))
        text = "\n".join(paras)
        print(json.dumps(dict(ok=True, paragraphs=len(paras), tables=len(d.tables), text=text[:14000], table_preview=tables[:3], note="%d characters" % len(text)), ensure_ascii=False))
    except Exception as e:
        print(json.dumps(dict(ok=False, error="read failed: %s" % str(e)[:200])))
        return 4
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
