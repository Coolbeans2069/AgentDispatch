"""probe an Excel workbook: sheet names, preview rows (openpyxl)."""
import json, os, sys


def main(argv):
    path = argv[0] if argv else ""
    sheet = argv[1] if len(argv) > 1 else ""
    if not path or not os.path.isfile(path):
        print(json.dumps(dict(ok=False, error="give an .xlsx/.xlsm file")))
        return 2
    try:
        import openpyxl
    except ImportError as e:
        print(json.dumps(dict(ok=False, error="openpyxl missing: %s" % e)))
        return 3
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        name = sheet if sheet in names else names[0]
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 25:
                break
            rows.append([("" if v is None else str(v)[:120]) for v in row[:40]])
        wb.close()
        print(json.dumps(dict(ok=True, sheets=names[:40], sheet=name, preview=rows, note="%d preview rows" % len(rows)), ensure_ascii=False))
    except Exception as e:
        print(json.dumps(dict(ok=False, error="read failed: %s" % str(e)[:200])))
        return 4
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
